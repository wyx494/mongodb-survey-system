# -*- coding: utf-8 -*-
"""生成《项目报告.docx》（叙述版）。"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(BASE)
SHOTS = os.path.join(ROOT, "docs", "screenshots")
OUT = os.environ.get("REPORT_OUT") or os.path.join(ROOT, "项目报告.docx")

doc = Document()

section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.left_margin = Cm(2.6)
section.right_margin = Cm(2.6)
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)

normal = doc.styles["Normal"]
normal.font.name = "宋体"
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(4)


def set_run(run, name="宋体", size=10.5, bold=False, color=None, italic=False):
    latin = "Times New Roman" if name == "宋体" else name
    run.font.name = latin
    rfonts = run._element.get_or_add_rPr().rFonts
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(text), "黑体", 16, True)


def h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(text), "黑体", 12.5, True)


def h3(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(text), "黑体", 11, True)


def para(text, indent=True, size=10.5, space_after=5):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    set_run(p.add_run(text), "宋体", size)
    return p


def bullets(items, size=10.5):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.line_spacing = 1.4
        p.paragraph_format.space_after = Pt(2)
        set_run(p.add_run("· " + item), "宋体", size)


def code(text, size=8.5):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.1
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F2F4F7")
    pPr.append(shd)
    run = p.add_run()
    for i, line in enumerate(text.split("\n")):
        if i:
            run.add_break()
        run.add_text(line)
    set_run(run, "Consolas", size)


def caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    set_run(p.add_run(text), "宋体", 9, color=(0x66, 0x66, 0x66))


def figure(filename, width=15.0, cap=None):
    path = os.path.join(SHOTS, filename)
    if os.path.exists(path):
        doc.add_picture(path, width=Cm(width))
        p = doc.paragraphs[-1]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
    if cap:
        caption(cap)


def shade_cell(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def table(headers, rows, widths=None, size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(str(h)), "黑体", size, True)
        shade_cell(hdr[i], "E8EEF7")
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            set_run(p.add_run(str(value)), "宋体", size)
    if widths:
        total = sum(widths)
        scale = 15.6 / total if total > 15.6 else 1.0
        widths = [w * scale for w in widths]
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def quote_box(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.right_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "FFF8E6")
    pPr.append(shd)
    set_run(p.add_run(text), "宋体", 9.5)


# 页脚页码
footer_p = section.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer_p.add_run("第 ")
set_run(run, "宋体", 9, color=(0x88, 0x88, 0x88))
run = footer_p.add_run()
fld1 = OxmlElement("w:fldChar")
fld1.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText")
instr.set(qn("xml:space"), "preserve")
instr.text = " PAGE "
fld2 = OxmlElement("w:fldChar")
fld2.set(qn("w:fldCharType"), "end")
run._r.append(fld1)
run._r.append(instr)
run._r.append(fld2)
set_run(run, "宋体", 9, color=(0x88, 0x88, 0x88))
run = footer_p.add_run(" 页")
set_run(run, "宋体", 9, color=(0x88, 0x88, 0x88))

# 标题
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(18)
set_run(p.add_run("基于 MongoDB 的在线问卷系统"), "黑体", 20, True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
set_run(p.add_run("第一阶段报告"), "黑体", 12, False, (0x55, 0x55, 0x55))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(2)
set_run(p.add_run("代码仓库：https://github.com/wyx494/mongodb-survey-system"), "宋体", 10.5, False, (0x33, 0x33, 0x33))

# ================= 一 =================
h1("一、需求分析与功能范围")

para(
    "作业的题目只有一句话：做一个简化版的在线问卷系统，功能类似问卷星。"
    "真开始想的时候，问题反而多了——“简化”到底简化到什么程度？哪些必须有，哪些可以不做？"
)

para(
    "我按使用者把需求过了一遍。出问卷的人要能登录、建问卷、加题、发链接、看结果；"
    "填问卷的人点开链接、答题、提交；还有一类是旁观者，比如老师，会看这个系统跑起来是什么样。"
    "三条线理清之后，功能范围基本就出来了。"
)

para(
    "有几处需求没有写死，我先停下来把它们定了。最绕的是匿名和登录的关系：文档一边说填写的人要先登录，"
    "一边又说问卷可以允许匿名。我的理解是，登录是门槛，匿名只是答卷里不记名，所以填写人仍然要登录，"
    "只是 respondentId 存 null。"
)

para(
    "重复填写也一样，做成问卷上的一个开关，默认允许。跳转我限定只能往后：发布前会校验目标题必须排在当前题后面，"
    "这样既不会出现死循环，也避免同一道题被反复回答。发布之后题目锁死，标题、说明和截止时间还能改——"
    "不然已经收上来的答卷和题目对不上，统计就没法解释了。"
)

para(
    "功能最后落到六块：账号、问卷、题型、跳转、填写、统计。做到什么程度，大概是下面这样："
)

table(
    ["模块", "做到什么程度"],
    [
        ["账号", "注册、登录、退出；密码 scrypt 加盐；只能看到自己的问卷"],
        ["问卷", "新建草稿、改标题/说明/匿名/重复填写/截止时间、发布、关闭、删草稿、复制短链"],
        ["题型", "单选、多选（可限制最少/最多/必须选几个）、文本填空（长度）、数字填空（范围、整数）"],
        ["跳转", "按单选、多选、数字的结果跳题；规则存在问卷里，可在编辑页配置"],
        ["填写", "登录后逐题作答、按规则跳、前后端都校验、匿名/记名、可限制重复"],
        ["统计", "单选计数、多选计数、填空列表、数字平均/最值"],
    ],
    widths=[2.6, 13.4],
)

para(
    "代码在 survey-system 目录下，数据库设计、接口和测试的说明放在 docs 里，测试用 npm test 跑。"
    "这份报告是最后写的，讲的是我怎么一步步把它做出来。"
)

# ================= 二 =================
h1("二、环境搭建与技术选型")

para(
    "环境是现成的：Windows 11，Node 22，MongoDB 8.3，VS Code。"
    "真正花时间的不是写代码，而是把 MongoDB 弄起来。"
)

para(
    "MongoDB 装是装了，但没初始化：没有注册服务，没有数据目录，配置文件里的 dbPath 写的是 %MONGO_DATA_PATH%，"
    "而这个变量根本没赋值。mongosh 连上去直接报 ECONNREFUSED。我一开始以为是端口被占，查了进程和端口才发现服务端根本没启动。"
)

para(
    "我没有重装，也没去改安装目录，直接在用户目录下建了数据目录，手动起服务："
)

code(
    'mkdir "%LOCALAPPDATA%\\MongoDB\\data"\n'
    '"C:\\Program Files\\MongoDB\\Server\\8.3\\bin\\mongod.exe" '
    '--dbpath "%LOCALAPPDATA%\\MongoDB\\data" --bind_ip 127.0.0.1 --port 27017'
)

para(
    "起来之后 ping、插入、计数、删库都试了一遍，确认读写正常。为了以后不用每次手敲，写了个 start-mongo.bat。"
    "数据放在用户目录而不是 C:\\data，主要是为了不需要管理员权限。"
)

para(
    "还有个小麻烦：PowerShell 默认禁止运行 npm.ps1，直接敲 npm install 会报“禁止运行脚本”。"
    "把终端切成 Command Prompt，或者用 npm.cmd 就正常了。VS Code 默认终端是 PowerShell，第一次装依赖时在这里卡了一会儿。"
)

para(
    "语言选 Node，因为机器上本来就有，写 HTTP 和 JSON 也顺。框架用 Express，没上 Nest 之类的东西——"
    "这个项目的路由不多，中间件也简单，Express 足够。"
)

para(
    "数据库驱动用的是官方 mongodb，没有用 Mongoose。这个选择有点刻意：Mongoose 会把 schema 藏进一层抽象里，"
    "而这次作业的重点恰恰是数据库设计。我想让集合、索引、校验器、聚合管道都明明白白写在代码里，"
    "讲的时候能直接指出来，而不是绕到某个 ORM 的 API 后面。"
)

para(
    "前端更简单：EJS 服务端渲染，加一点原生 JS。需求明说不要求复杂界面，引入 React 反而多一层构建和一堆依赖。"
    "项目结构按“入口、配置、数据、路由、服务、页面”分开，后面加功能时知道该改哪："
)

code(
    "survey-system/\n"
    "├─ server.js            启动入口：连库、建索引、监听端口\n"
    "├─ src/\n"
    "│  ├─ app.js            Express 组装（中间件、路由挂载、错误处理）\n"
    "│  ├─ config.js         端口、连接串、库名、会话时长\n"
    "│  ├─ db.js             连接、集合校验器、索引\n"
    "│  ├─ auth.js           密码哈希、会话、登录中间件\n"
    "│  ├─ routes/           auth / surveys / public / stats / pages\n"
    "│  ├─ services/         jump（跳转）、validate（校验）、stats（聚合）、survey（结构）\n"
    "│  └─ views/            EJS 页面\n"
    "├─ public/              样式和前端脚本\n"
    "├─ scripts/             初始化、演示数据、检查、截图\n"
    "├─ tests/               单元测试 + API 集成测试\n"
    "└─ docs/                数据库设计、API、测试用例、AI 日志、截图"
)

# ================= 三 =================
h1("三、数据库设计")

para(
    "这是花时间最多的一部分。动笔之前我先拿纸画了几种方案：所有东西塞一个集合、题目单独一个集合、答卷内嵌进问卷……"
    "最后落到现在这四个集合。"
)

table(
    ["集合", "存什么", "为什么单独放"],
    [
        ["users", "账号、密码哈希、注册时间", "登录要按用户名查，和问卷数据没关系"],
        ["surveys", "问卷信息 + 题目数组（内嵌）", "题目和问卷总是一起读，数量也有上限"],
        ["responses", "每一份答卷", "会持续增长，写入频繁，还要单独聚合"],
        ["sessions", "登录会话", "读写频繁，靠 TTL 自动过期"],
    ],
    widths=[2.4, 7.0, 6.6],
)

h2("3.1 题目为什么内嵌")

para(
    "编辑一份问卷时，我要一次拿到所有题目和选项；发布校验要遍历所有跳转；填写页也要整份题目。"
    "如果拆成单独的 questions 集合，每次都要多查一次再拼起来，改一道题还要跨文档更新，既慢又容易出中间状态。"
)

para(
    "题目数量有上限（我限制最多 100 题），内嵌不会把文档撑爆。跳转规则也跟着题目走——它本来就是这道题的属性，"
    "“选 A 去第 5 题”离开这道题就没有意义。这样发布校验、路径计算、编辑器渲染都只需要遍历一个数组。"
)

h2("3.2 答卷为什么单独存")

para(
    "答卷的情况正好相反。一份问卷可能收几百上千份，如果内嵌，同一个文档会被反复写，16MB 上限迟早撞上；"
    "多人同时提交时还会互相锁。拆开之后，每次提交只是插一个新文档，互不干扰。统计也只扫 responses，"
    "不用把整份问卷读出来。"
)

para(
    "我在答卷里多加了一个 path 字段，记录这次实际走了哪些题。需求没要求这个字段，但测试和排查都靠它——"
    "比如断言第 1 题选 D 之后，path 就是 [\"q1\",\"q4\"]。"
)

h2("3.3 索引")

para(
    "索引是按查询场景一条条对的，不是把每个字段都建一遍。"
)

table(
    ["集合", "索引", "对应什么操作"],
    [
        ["users", "{ usernameLower: 1 } unique", "登录、注册查重，大小写不敏感"],
        ["surveys", "{ slug: 1 } unique", "短链访问"],
        ["surveys", "{ ownerId: 1, createdAt: -1 }", "我的问卷列表"],
        ["surveys", "{ status: 1 }", "按状态筛选"],
        ["responses", "{ surveyId: 1, submittedAt: -1 }", "按问卷取答卷、按时间排序"],
        ["responses", "{ surveyId: 1, respondentId: 1 }", "判断这个人是不是已经填过"],
        ["sessions", "{ expiresAt: 1 } TTL", "过期会话自动清理"],
        ["sessions", "{ userId: 1 }", "按用户清理，暂时没用上"],
    ],
    widths=[2.4, 7.6, 6.0],
)

para(
    "用户名为什么要多存一个 usernameLower？因为 MongoDB 的唯一索引区分大小写，不存小写副本的话，"
    "Alice 和 alice 会变成两个账号。用副本加唯一索引，比查询时写 $regex 干净得多。"
)

para(
    "会话没有放在 Node 内存里，而是写进 sessions 集合。好处是服务重启不会把所有人踢下线；"
    "expiresAt 上建了 TTL 索引，过期文档由 MongoDB 自己删，不用写定时任务。"
)

h2("3.4 为什么不用关系型")

para(
    "如果硬套关系型，大概会拆成 users、surveys、questions、options、responses、answers 六张表。"
    "一次填写要跨三四张表关联，统计要写多层 JOIN 加 GROUP BY。"
)

para(
    "这个项目的形状其实很适合文档模型：题型异构（单选、多选、文本、数字，字段不一样）、题目和选项天然嵌套、"
    "答案也是按题组织的嵌套结构、统计要按题分组再按选项展开。"
    "一份问卷就是一个文档，一次填写就是一份答卷文档，统计直接对数组做 $unwind 和 $group。"
)

para(
    "另外，第二阶段还要改需求。文档模型改字段比改表结构轻，新增题型不用动已有数据。"
    "这也是我选 MongoDB 的原因之一，不只是因为课程要求。"
)

h2("3.5 文档长什么样")

para("问卷文档（省略了时间字段）：")

code(
    '{\n'
    '  "ownerId": ObjectId("..."),\n'
    '  "slug": "YFttks6f",\n'
    '  "title": "课程学习体验调查",\n'
    '  "anonymous": false,\n'
    '  "allowMultiple": true,\n'
    '  "status": "published",\n'
    '  "questions": [\n'
    '    {\n'
    '      "qid": "q1", "type": "single", "required": true,\n'
    '      "title": "你现在读大几？",\n'
    '      "options": [{"key": "A", "label": "大一"}, {"key": "B", "label": "大二"}],\n'
    '      "constraints": {},\n'
    '      "jumps": [{"when": {"option": "D"}, "to": "q4"}]\n'
    '    }\n'
    '  ]\n'
    '}'
)

para("答卷文档：")

code(
    '{\n'
    '  "surveyId": ObjectId("..."),\n'
    '  "respondentId": ObjectId("...") 或 null（匿名）,\n'
    '  "anonymous": false,\n'
    '  "answers": [\n'
    '    {"qid": "q1", "value": "B"},\n'
    '    {"qid": "q2", "value": ["A", "C"]},\n'
    '    {"qid": "q3", "value": 8}\n'
    '  ],\n'
    '  "path": ["q1", "q2", "q3"],\n'
    '  "submittedAt": ISODate("...")\n'
    '}'
)

para(
    "除了应用层校验，四个集合都加了 $jsonSchema 校验器：字段类型、必填项、status 的枚举值。"
    "比如 surveys.status 只能是 draft/published/closed。这样即使有人绕过接口直接写库，也不会写出结构错误的数据。"
)

para("几条一致性上的约定，后面写代码时一直按这个来：")

table(
    ["场景", "怎么处理", "为什么"],
    [
        ["发布后改题目", "禁止，返回 409", "否则已收答卷和题目对不上，统计会失真"],
        ["匿名问卷", "仍要登录，respondentId 存 null", "登录是门槛，匿名只是不记名"],
        ["重复填写", "问卷级开关 allowMultiple", "需求里说的是“如果允许”"],
        ["截止时间", "提交时校验", "到点后拒绝新答卷，已收的不受影响"],
        ["跳转", "只能向后，发布前校验", "防环，也避免同一题被反复作答"],
    ],
    widths=[3.2, 6.4, 6.0],
)

# ================= 四 =================
h1("四、后端实现")

para(
    "后端是按依赖顺序长的：先让服务能连库、能注册登录，再做问卷，最后做填写和统计。"
    "每做完一块就用 curl 或测试打一遍，不攒到最后一起调。"
)

h2("4.1 登录和会话")

para(
    "密码用 Node 内置的 scryptSync 加盐哈希：每个用户生成 16 字节随机盐，scrypt 派生 64 字节，只存哈希和盐。"
    "登录时重新派生，用 timingSafeEqual 比较，避免时序侧信道。"
)

para(
    "会话写进 sessions 集合：登录时生成 32 字节随机 token 当 _id，同时写 userId 和 expiresAt；"
    "Cookie 里只放这个 token（HttpOnly、SameSite=Lax）。每次请求按 token 查会话，再查用户。"
    "这样服务重启不会掉登录，多进程也能共享。"
)

h2("4.2 问卷和发布校验")

para(
    "创建问卷时生成 8 位短码，去掉了 0/o、1/l/I 这些容易看错的字符；撞了唯一索引就重试，最多 5 次。"
    "刚创建是草稿，题目为空。编辑接口对字段做了白名单，标题、说明、匿名、重复填写、截止时间、题目分别处理。"
)

para(
    "题目只有草稿能改，已发布的问卷改题目会返回 409。标题和截止时间不受限制，因为改这些不会让历史答卷对不上。"
    "发布前会检查：至少一道题、题号不重复、选择题至少两个选项、选项 key 不重复、跳转目标存在且排在当前题后面。"
    "任何一条不过都不让发布，并告诉具体是哪一题。"
)

h2("4.3 跳转")

para(
    "每条规则大概长这样：单选 {\"when\": {\"option\": \"A\"}, \"to\": \"q5\"}，"
    "多选 {\"when\": {\"includes\": \"A\"}, \"to\": \"q3\"}，"
    "数字 {\"when\": {\"op\": \"gte\", \"value\": 60}, \"to\": \"q6\"}。"
    "按顺序匹配，命中就跳，没命中走下一题。"
)

para("核心就是下面这个循环：")

code(
    "function computePath(questions, answers) {\n"
    "  const indexById = new Map(questions.map((q, i) => [q.qid, i]));\n"
    "  const path = [];\n"
    "  const visited = new Set();\n"
    "  let index = 0;\n"
    "  while (index >= 0 && index < questions.length) {\n"
    "    const question = questions[index];\n"
    "    if (visited.has(question.qid)) throw new Error(`检测到循环跳转：${question.qid}`);\n"
    "    visited.add(question.qid);\n"
    "    path.push(question.qid);\n"
    "    const answer = answers.get(question.qid);\n"
    "    const rule = (question.jumps || []).find((r) => matchRule(r, answer));\n"
    "    if (rule) {\n"
    "      const target = indexById.get(rule.to);\n"
    "      if (target === undefined) throw new Error(`跳转目标不存在：${rule.to}`);\n"
    "      index = target;\n"
    "    } else {\n"
    "      index += 1;\n"
    "    }\n"
    "  }\n"
    "  return path;\n"
    "}"
)

para(
    "visited 是防止数据被改坏时死循环；更根本的保证在发布校验：目标必须在当前题后面，正常情况下不可能成环。"
    "两道防线，一道防运行时，一道防数据。"
)

h2("4.4 校验")

para(
    "填写页在前端也会算跳转，但服务端不能信前端。提交时服务端把答案放进 Map，先算出这次实际经过的题，"
    "再只校验路径上的题目。被跳过的题即使带了答案也不会校验，更不会存库。"
)

para(
    "单选要求必答且选项存在；多选要满足最少、最多或必须选几个；文本看长度，存库前 trim；"
    "数字看范围、是否整数。答案里出现问卷没有的题号直接拒绝，避免脏数据。"
)

h2("4.5 统计")

para(
    "统计全部在 MongoDB 里算。单选题的管道是：match 出这份问卷、unwind 答案、match 到这道题、"
    "按答案值 group 计数。多选题多一层 unwind，因为答案本身是数组。数字题一次分组算 count、avg、min、max。"
)

code(
    "db.responses.aggregate([\n"
    '  { $match: { surveyId } },\n'
    '  { $unwind: "$answers" },\n'
    '  { $match: { "answers.qid": "q1" } },\n'
    '  { $group: { _id: "$answers.value", count: { $sum: 1 } } }\n'
    "])"
)

para(
    "这里有个统计口径的问题，后面第七节会细说：作答人数不能把选项次数加起来。"
    "多选题一个人选 3 项，选项次数会贡献 3 次，求和就把人数算多了。"
)

h2("4.6 接口")

table(
    ["方法", "路径", "作用"],
    [
        ["POST", "/api/auth/register", "注册"],
        ["POST", "/api/auth/login", "登录"],
        ["POST", "/api/auth/logout", "退出"],
        ["GET", "/api/auth/me", "当前用户"],
        ["GET", "/api/surveys", "我的问卷列表"],
        ["POST", "/api/surveys", "新建草稿"],
        ["GET", "/api/surveys/:id", "问卷详情"],
        ["PATCH", "/api/surveys/:id", "修改问卷/题目"],
        ["POST", "/api/surveys/:id/publish", "发布"],
        ["POST", "/api/surveys/:id/close", "关闭"],
        ["DELETE", "/api/surveys/:id", "删除草稿"],
        ["GET", "/api/public/surveys/:slug", "填写页取卷"],
        ["POST", "/api/public/surveys/:slug/responses", "提交答卷"],
        ["GET", "/api/surveys/:id/stats", "整卷统计"],
        ["GET", "/api/surveys/:id/stats/:qid", "单题统计"],
    ],
    widths=[2.0, 9.2, 5.8],
)

# ================= 五 =================
h1("五、前端实现")

para(
    "页面一共五个：登录/注册、我的问卷、编辑问卷、填写页、统计页。样式是一套很普通的 CSS，"
    "没有用组件库，也没有做移动端适配。"
)

para(
    "编辑器是最费事的一块。题目放在一个数组里，输入时只改数组、不重绘，避免输入框失焦；"
    "只有增删题目、增删选项、增删跳转、切换题型时才整体重绘。全量重绘简单，但每敲一个字就重绘会让光标跳走，"
    "所以最后选了这种折中的做法。"
)

para(
    "跳转规则的目标题下拉框只列出当前题后面的题，从源头上避免往前跳。发布之后，题目区域整体禁用，"
    "页面上会提示“发布后题目不能再修改”。"
)

figure("03-edit.png", 14.5, "图 1　编辑器顶部：问卷信息、匿名与重复填写开关、截止时间")
figure("03b-edit-jumps.png", 14.5, "图 2　同一页往下：选项、数量限制和跳转规则")

para(
    "填写页一次只显示一道题，顶部有进度条。点“下一题”时先做本地校验，通过后按规则决定下一题；"
    "如果下一题越界，就提交。前端算跳转是为了即时反馈，但真正算数的是服务端——提交时服务端会重新算一遍路径。"
)

figure("04-fill-q1.png", 12.5, "图 3　填写页：一次一题，进度条显示当前位置")
figure("05-fill-q4.png", 12.5, "图 4　第 1 题选“大四”后直接跳到第 4 题")

para(
    "统计页是服务端渲染的：选择题用表格加条形长度表示占比，填空题列出全部文本，数字题显示平均、最小、最大。"
    "数据来自 computeStats，和接口用的是同一个函数，避免页面和接口两套口径。"
)

figure("06-stats.png", 14.5, "图 5　统计页：单选计数、多选计数和占比")

# ================= 六 =================
h1("六、测试与结果")

para(
    "测试分两层。单元测试覆盖跳转、校验、题目结构这些纯函数；API 测试起一个真实的 Express 实例，"
    "连独立的测试库，用 fetch 打接口，覆盖注册登录、建卷、发布、填写、统计、越权、重复填写、关闭和截止。"
    "用 Node 自带的 node:test，不另外装 Jest——接口测试本来就不需要浏览器。"
)

table(
    ["层级", "文件", "数量", "覆盖"],
    [
        ["单元测试", "tests/unit.test.js", "14 个 test", "matchRule、computePath、validateSubmission、normalizeQuestions、validateForPublish、generateSlug"],
        ["API 集成", "tests/api.test.js", "23 个子测试", "完整业务链路和异常分支"],
        ["浏览器检查", "scripts/ui-check.js", "2 条链路", "编辑器保存、完整填写提交"],
    ],
    widths=[2.6, 5.2, 4.0, 8.0],
)

para(
    "测试库用 survey_test_时间戳，每次新建、跑完 dropDatabase，不碰演示数据。跑起来只有一条命令："
)

code("npm test\n\n# tests 37\n# pass 37\n# fail 0")

para(
    "第一轮其实挂了 5 个。三个是测试自己写错了：有的用例没给答案，跳转规则根本不会触发，却断言应该抛错；"
    "有的文本长度不够，规则要求 5 个字，测试只写了 4 个；还有查询 _id 时传了字符串，MongoDB 不会自动转 ObjectId。"
    "另外两个是真问题：多选题的人数算多了，以及 npm test 里 node --test tests/ 在 Windows 下把目录当成了模块。"
)

para("界面上的东西没法全用断言覆盖，所以又手工走了一遍：")

table(
    ["场景", "操作", "预期", "实际"],
    [
        ["重复用户名", "注册 demo", "提示用户名已被占用", "符合"],
        ["发布空问卷", "不添加题目直接发布", "提示至少一道题", "符合"],
        ["必答提示", "跳过必答题", "提示这道题是必答题", "符合"],
        ["多选上限", "选 3 个，上限 2", "提示最多选 2 个", "符合"],
        ["数字范围", "年龄填 200", "提示不能大于 120", "符合"],
        ["跳转", "第 1 题选大四", "直接显示第 4 题", "符合"],
        ["复制链接", "点复制链接", "得到 /survey/xxxx", "符合"],
    ],
    widths=[3.2, 6.4, 6.2, 2.0],
)

para(
    "最后还用 Playwright 打开真实浏览器走了一遍：登录、编辑保存、第 1 题选 D 跳到第 4 题、完整提交、看统计。"
    "这一步主要是确认前端脚本没有在真实浏览器里出问题——单元测试和接口测试都测不到这些。"
)

# ================= 七 =================
h1("七、遇到的问题与解决")

h3("统计里“人数”被算多了")
para(
    "统计测试期望是 2，实际跑出来 4。查了半天，发现多选题的“作答人数”是按选项次数求和的："
    "一个人选了 A、B 两项，就贡献了两次，人数自然翻倍。后来人数单独用 countDocuments 查，选项次数才用聚合。"
    "这个错误很隐蔽，因为单选题不会暴露——单选一个人只贡献一次。"
)

h3("路由顺序：/stats 被 /:id 吃掉")
para(
    "访问 /api/surveys/:id/stats 时，返回的是问卷详情。原因是 surveys 路由里的 GET /:id 先注册，"
    "把 /:id/stats 当成了 id 的一部分。把统计路由挂到问卷路由前面就好了。这种问题不报错，只是静默返回错的东西，"
    "不看返回内容很难发现。"
)

h3("测试挂了，先怀疑代码还是先怀疑测试")
para(
    "“跳转目标不存在应报错”的用例一直不抛错。看了半天代码，最后发现是用例没给答案——没有答案时规则不会命中，"
    "自然不会走到目标。这件事提醒我，测试挂了不一定是代码错，也可能是测试的前提没构造对。"
    "后来我养成了先看测试在测什么、再怀疑实现的习惯。"
)

h3("编辑器里的两个小毛病")
para(
    "跳转目标的下拉框保存后重新打开不回显，显示的是“选择题目”，但数据库里的规则其实还在——"
    "生成 option 时没给当前目标加 selected。另一个是整页截图时底部的保存按钮出现在第一题中间，"
    "原因是操作栏用了 position: sticky。一个改 selected，一个改成普通布局，都不难，但不看真实页面发现不了。"
)

h3("匿名和登录的别扭")
para(
    "需求里这两句话放一起是矛盾的。我一开始想干脆让匿名问卷免登录，后来觉得不对："
    "文档明确说“填写的人需要先登录”，匿名更可能是指不记录身份。最后定成登录照旧、respondentId 存 null，"
    "并在 README 和报告里都写清楚。这个决定没有标准答案，但必须前后一致。"
)

h3("截图脚本选错了问卷")
para(
    "加了一份草稿做演示之后，脚本取列表第一份问卷，正好取到了草稿，填写页就 404 了。"
    "改成按状态取“已发布”的那份。顺带说，自动化脚本里这种“取第一个”的写法很脆弱，数据一变就翻车。"
)

# ================= 八 =================
h1("八、AI 使用说明")

para(
    "这个项目从需求梳理、写代码、跑测试到排版报告，基本每个环节都用了 AI。"
    "工具是 DeepSeek Harness 里的对话式助手，能读写本地文件、执行命令，所以我不是只问它问题，"
    "也让它直接改代码、跑命令。"
)

para(
    "最开始我让它把作业文档拆成功能清单和验收标准，并要求它单独列出文档里没写死的点。"
    "它列了匿名和登录的关系、重复填写、跳转方向、发布后能否改题、截止之后的行为这几条。"
    "后面我按课程要求逐条定了：匿名只是不记名、登录照旧；重复填写做成开关；跳转只往后；"
    "发布后锁题；截止后拒绝提交。"
)

para(
    "数据模型是我让它先出方案的重点。我让它设计 MongoDB 的集合结构，说明哪些字段内嵌、哪些用引用，"
    "并解释为什么不用关系型。它给了 users、surveys、responses 三个集合：题目内嵌进问卷，答卷单独存。"
    "我在此基础上把 sessions 拆出来单独放，加了 TTL 索引，又在答卷里补了 path 字段，记录实际作答路径。"
)

para(
    "跳转引擎也是我描述规则、它写初版。我告诉它：题目按顺序排，每道题可以挂若干条规则，"
    "单选按选项、多选按包含、数字按比较，命中就跳，没命中走下一题。它给了 matchRule 和 computePath。"
    "我检查时发现两个问题：没有防循环，目标题是否存在也没校验。后来让它加了 visited 集合和报错，"
    "又在发布校验里限定只能往后跳。"
)

para(
    "答案校验和统计也是我先说清楚口径再让它写。校验方面我要求只校验实际路径上的题，"
    "它给了 validateSubmission 的骨架；我补了未知题号直接拒绝、文本 trim。"
    "统计方面我让它用聚合管道写，它把单选、多选、文本、数字四类都写出来了，"
    "但多选题的“作答人数”是拿选项次数相加——一个人选三项就算成三个人。"
    "我改成单独 countDocuments 统计人数，选项次数才用聚合。"
)

para(
    "测试用例是我让它按模块列，覆盖注册登录、建卷、发布、跳转、校验、提交、统计、越权、重复填写、关闭和截止。"
    "它一次给了二十多个。第一轮跑挂了 5 个：其中 3 个是它用例本身写错——前提没构造、文本长度不够、"
    "ObjectId 传了字符串；另外 2 个是代码问题——多选人数算多了，以及 npm test 的目录参数在 Windows 下不对。"
)

para(
    "前端编辑器也交给它写。我要求不用框架、原生 JS，能增删题目和选项、配置跳转、保存成 JSON。"
    "它给了“状态数组 + 结构变化时重绘、输入时只改数组”的结构，避免输入框失焦。"
    "后面发现跳转目标的下拉框保存后不回显，我补了 selected；底部操作栏用了 sticky，整页截图时会盖住内容，"
    "改成了普通布局。"
)

para(
    "报告和截图脚本也用了 AI。我让它按“需求—设计—实现—测试—问题—AI”列提纲，再逐节扩写；"
    "截图脚本让它用 Playwright 打开页面、截图，并验证第 1 题选 D 会跳到第 4 题。"
    "脚本第一版取列表里第一份问卷，后来加了一份草稿就取错了，我改成按状态取“已发布”的那份。"
)

para(
    "它出错的地方集中在几类：统计口径、跳转防循环、测试前提、编辑器回显、脚本取错对象。"
    "这些不是语法错误，而是它不知道课程约定和业务口径——它不知道匿名还要登录，"
    "也不知道多选的人数不能按选项次数算。"
)

para(
    "所以有些决定我没有让它做：四个集合怎么划分、匿名和登录怎么解释、发布后要不要锁题、"
    "跳转能不能往前、“人数”按什么算、校验的边界。它给的是通用实现，这些取舍要按课程要求自己定。"
    "报告里的设计和问题记录，也是按这些实际决定写的。"
)

# ================= 九 =================
h1("九、第二阶段的准备")

para(
    "第二阶段会发新的需求变更，而且要求在这个系统上改、不能推倒重做。所以我这一版刻意留了几个地方："
    "问卷文档带 schemaVersion，将来结构变化可以做迁移；题型用字符串枚举，新增题型只需要加分支，"
    "不用动历史数据；跳转规则是数组，以后要支持组合条件，可以往 when 里加字段，旧规则仍然有效；"
    "统计按题目循环生成，新增题型补一个聚合分支就行。"
)

para(
    "最可能要动的是“发布后锁题”这个决定。如果新需求要求已发布的问卷也能改题目，就得引入版本化："
    "一份问卷多个版本，答卷记录当时答的是哪个版本。其次是协作编辑或题目分组——现在的 questions 是一个扁平数组，"
    "ownerId 也只有一个，真要多人协作，结构得调整。好在现在数据量不大，迁移成本可控。"
)

# ================= 附录 =================
h1("附录 A　数据库字段速查")

table(
    ["集合", "字段", "类型", "说明"],
    [
        ["users", "username / usernameLower", "string", "显示名 / 小写唯一键"],
        ["users", "passwordHash / passwordSalt", "string", "scrypt 哈希与盐"],
        ["users", "createdAt", "date", "注册时间"],
        ["surveys", "ownerId", "ObjectId", "创建者"],
        ["surveys", "slug", "string", "8 位短码，唯一"],
        ["surveys", "status", "string", "draft / published / closed"],
        ["surveys", "anonymous / allowMultiple", "bool", "匿名 / 允许重复"],
        ["surveys", "deadline", "date | null", "截止时间"],
        ["surveys", "questions", "array", "题目、选项、约束、跳转"],
        ["responses", "surveyId", "ObjectId", "所属问卷"],
        ["responses", "respondentId", "ObjectId | null", "填写人，匿名时 null"],
        ["responses", "answers", "array", "只含实际经过的题"],
        ["responses", "path", "array", "实际作答路径"],
        ["responses", "submittedAt", "date", "提交时间"],
        ["sessions", "_id", "string", "会话 token"],
        ["sessions", "userId / expiresAt", "ObjectId / date", "用户与过期时间（TTL）"],
    ],
    widths=[2.2, 5.8, 3.8, 5.2],
)

h1("附录 B　接口速查")

table(
    ["方法", "路径", "说明"],
    [
        ["POST", "/api/auth/register", "注册并登录"],
        ["POST", "/api/auth/login", "登录"],
        ["POST", "/api/auth/logout", "退出"],
        ["GET", "/api/auth/me", "当前用户"],
        ["GET / POST", "/api/surveys", "列表 / 新建"],
        ["GET / PATCH / DELETE", "/api/surveys/:id", "详情 / 修改 / 删除草稿"],
        ["POST", "/api/surveys/:id/publish", "发布"],
        ["POST", "/api/surveys/:id/close", "关闭"],
        ["GET", "/api/public/surveys/:slug", "公开取卷"],
        ["POST", "/api/public/surveys/:slug/responses", "提交答卷"],
        ["GET", "/api/surveys/:id/stats", "整卷统计"],
        ["GET", "/api/surveys/:id/stats/:qid", "单题统计"],
    ],
    widths=[4.4, 8.0, 4.6],
)

doc.save(OUT)
print("saved:", OUT)
